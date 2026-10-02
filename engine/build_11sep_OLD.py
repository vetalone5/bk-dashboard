# -*- coding: utf-8 -*-
import openpyxl, json, re, datetime
U='../data/tables/'
EL=U+'01_elena_influence.xlsx'
MAR=U+'02_maria_influence.xlsm'
NAS=U+'03_anastasia_marketplace.xlsx'
ANYA=U+'04_anya_external_and_bot.xlsx'

MONTHS={'2026-03':{'label':'Март 2026','weeks':['1–7 мар','8–14 мар','15–21 мар','22–28 мар','29–31 мар']},
        '2026-04':{'label':'Апрель 2026','weeks':['1–7 апр','8–14 апр','15–21 апр','22–28 апр','29–30 апр']},
        '2026-05':{'label':'Май 2026','weeks':['1–7 май','8–14 май','15–21 май','22–28 май','29–31 май']},
        '2026-06':{'label':'Июнь 2026','weeks':['1–7 июн','8–14 июн','15–21 июн','22–28 июн','29–30 июн']},
        '2026-07':{'label':'Июль 2026','weeks':['1–7 июл','8–14 июл','15–21 июл','22–28 июл','29–31 июл']},
        '2026-08':{'label':'Август 2026','weeks':['1–9 авг','10–16 авг','17–23 авг','24–30 авг','31 авг']},
        '2026-09':{'label':'Сентябрь 2026','weeks':['1–6 сен','7–13 сен','14–20 сен','21–27 сен','28–30 сен']}}
def widx(d):
    dd=d.day
    if d.month==8: return 0 if dd<=9 else 1 if dd<=16 else 2 if dd<=23 else 3 if dd<=30 else 4
    if d.month==9: return 0 if dd<=6 else 1 if dd<=13 else 2 if dd<=20 else 3 if dd<=27 else 4
    return 0 if dd<=7 else 1 if dd<=14 else 2 if dd<=21 else 3 if dd<=28 else 4
def num(x):
    if x is None: return 0
    if isinstance(x,(int,float)): return float(x)
    s=re.sub(r'[^0-9.\-]','',str(x).replace('\xa0','').replace(' ','').replace(',','.'))
    try: return float(s) if s not in('','-','.') else 0
    except: return 0
def txt(x): return ('' if x is None else str(x)).strip()
def blank(): return [[0,0,0,0,0,0,0,0] for _ in range(5)]
CH={'max':'MAX','макс':'MAX','vk':'VK','vkontakte':'VK','вк':'VK','inst':'Instagram','instagram':'Instagram',
    'youtube':'YouTube','telegram':'Telegram','tiktok':'TikTok','rutube':'Rutube','тг':'Telegram',
    'ютуб':'YouTube','тик-ток':'TikTok','vk клипы':'VK','vk видео':'VK'}
def chan(s):
    k=txt(s).lower().strip()
    if k in ('wb','вб','ozon','озон','сайт',''): return 'Канал не указан'
    return CH.get(k, txt(s) or 'Канал не указан')
def plat(s):
    k=txt(s).lower()
    if 'wb' in k or 'вб' in k: return 'Wildberries'
    if 'ozon' in k or 'озон' in k: return 'Ozon'
    if 'сайт' in k: return 'Сайт trigger-fish.ru'
    return 'Не указано'
def platByTag(art,utm,link,promo):
    a=txt(art).upper(); u=(txt(utm)+' '+txt(link)).lower()
    if a.startswith('WW') or 'wildberries' in u: return 'Wildberries'
    if 'ozon' in u: return 'Ozon'
    if 'trigger-fish' in u: return 'Сайт trigger-fish.ru'
    return 'Не указано'
def tagfor(pl,art,utm,promo):
    art,utm,promo=txt(art),txt(utm),txt(promo)
    if promo in ('нет','-','—','-TRG'): promo=''
    if art in ('нет','-','—'): art=''
    if pl.startswith('Сайт'): return 'utm' if utm else ('promo' if promo else 'none')
    if pl=='Wildberries': return 'art' if art.upper().startswith('WW') else ('promo' if promo else 'none')
    if pl=='Ozon': return 'utm' if 'utm_' in utm else ('promo' if promo else 'none')
    return 'none'
BOTS={'gd':'Бот «Зелёный доктор»','us':'Бот «Умный сад»','km':'Бот «Клёвое место»'}
def prodof(x):
    k=txt(x).lower()
    if 'зелён' in k or 'зелен' in k or 'доктор' in k or 'green_doctor' in k: return 'gd'
    if 'умный сад' in k or 'умная гряд' in k or 'garden_app' in k: return 'us'
    if 'клёв' in k or 'клев' in k: return 'km'
    if 'апплик' in k: return 'appl'
    return 'atf'
def done(x):
    return txt(x).lower().startswith(('вышл','опублик'))
PATS=[('spend','итого к оплате'),('osum','сумма заказов'),('orders','заказы'),('buy','выкупы'),('rev','выручка от выкупов'),
      ('clm','клики мобзио'),('clw','клики wb'),('reach','охват'),('cart','корзины'),
      ('date','дата выхода'),('status','статус'),('plat','площадка'),('sales','канал продаж'),
      ('art','подменный артикул'),('utm','ссылка с utm'),('promo','промокод'),
      ('link','ссылка на размещение'),('prod','продукт'),('scn','сценарий'),('reg','регистрации')]
def cmap(hdr):
    m={}
    for i,c in enumerate(hdr):
        t=txt(c).lower().replace('\n',' ')
        for k,p in PATS:
            if t.startswith(p) and k not in m: m[k]=i
    return m

# ---------- ДАННЫЕ ИЗ API (папка api/) ----------
import os
def loadapi(n):
    try: return json.load(open('../data/api/'+n))
    except Exception as e: print('api',n,'нет:',e); return {}
A_UTM=loadapi('ozon_utm.json'); A_OZADS=loadapi('ozon_ads.json')
A_WBADS=loadapi('wb_ads.json'); A_WBFUN=loadapi('wb_funnel.json')
A_OZORD=loadapi('ozon_orders.json'); A_RATE=loadapi('rating.json')
A_SHORT=loadapi('shortlinks.json')
A_DEEP=loadapi('ozon_utm_deep.json')
GENERIC={'vendor_org_4202727'}
def urlof(u,link=None):
    u=txt(u)
    m=re.search(r'onelink\.me/SNMZ/([A-Za-z0-9]+)',u)
    if m and A_SHORT.get(m.group(1)):
        return A_SHORT['full'].get(m.group(1),u) if isinstance(A_SHORT.get('full'),dict) else u
    return u
def parts(u):
    u=txt(u)
    if not u: return {}
    m=re.search(r'onelink\.me/SNMZ/([A-Za-z0-9]+)',u)
    if m:
        c=A_SHORT.get(m.group(1))
        return {'camp':c} if c else {}
    d={}
    for k,pat in (('camp',r'utm_campaign=([^&\s]+)'),('cont',r'utm_content=([^&\s]+)'),
                  ('term',r'utm_term=([^&\s]+)'),('src',r'utm_source=([^&\s]+)')):
        mm=re.search(pat,u)
        if mm: d[k]=mm.group(1).replace('+',' ').replace('%0A','').strip()
    mm=re.search(r'[?&]c=([^&\n]+)',u)
    if mm and 'camp' not in d: d['camp']=mm.group(1).strip()
    return d
def campof(u):
    u=txt(u)
    if not u: return None
    m=re.search(r'onelink\.me/SNMZ/([A-Za-z0-9]+)',u)
    if m: return A_SHORT.get(m.group(1))
    m=re.search(r'[?&]c=([^&\n]+)',u)
    if m: return m.group(1).strip()
    m=re.search(r'utm_campaign=([^&\s]+)',u)
    if m:
        c=m.group(1).replace('+',' ')
        return c if c.startswith('vendor_org') else None
    return None

R=[]; PEND={}
def scan(path,sheets,mgr):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    tot={}
    for sh,MM in sheets:
        ws=wb[sh]; rows=list(ws.iter_rows(min_row=5,values_only=True))
        if not rows: continue
        C=cmap(rows[0]); mon=int(MM[-2:])
        for r in rows[1:]:
            if not r or not r[0]: continue
            nm=txt(r[0])
            if nm.startswith(('НЕДЕЛЯ','Итого','ИТОГО')): continue
            g=lambda k: r[C[k]] if k in C and C[k]<len(r) else None
            if not done(g('status')):
                p=PEND.setdefault((MM,mgr),[0,0]); p[0]+=1; p[1]+=num(g('spend')); continue
            d=g('date')
            if not isinstance(d,datetime.datetime):
                m=re.match(r'(\d{1,2})[/.](\d{1,2})',txt(d))
                d=datetime.datetime(2026,mon,int(m.group(1))) if m else datetime.datetime(2026,mon,15)
            if d.month!=mon: d=datetime.datetime(2026,mon,min(d.day,28))
            pl=plat(g('sales'))
            if pl=='Не указано': pl=platByTag(g('art'),g('utm'),g('link'),g('promo'))
            sp=num(g('spend'))
            prod=prodof(g('prod'))
            w=blank(); w[widx(d)]=[sp,num(g('orders')),num(g('buy')),num(g('rev')),
                                   num(g('clm')) or num(g('clw')),1,num(g('osum')),num(g('reg'))]
            tg=tagfor(pl,g('art'),g('utm'),g('promo'))
            P=parts(g('utm')) or parts(g('link'))
            cmp_=P.get('camp'); ckey=None
            if cmp_ and cmp_ not in GENERIC: ckey=('camp',cmp_)
            elif P.get('cont'): ckey=('cont',P['cont'])
            elif P.get('term'): ckey=('term',P['term'])
            elif cmp_: ckey=('camp',cmp_)
            sc=txt(g('scn'))
            if sc.endswith('.0'): sc=sc[:-2]
            if sc in ('-','—','нет'): sc=''
            R.append(dict(m=MM,prod=prod,plat=pl,chan=chan(g('plat')),act='Инфлюенс-посев',name=nm[:46],
                link=txt(g('link')) or None,date=d.strftime('%d.%m'),tag=tg,mgr=mgr,w=w,est=tg=='none',
                scn=sc or None,camp=cmp_,ckey=list(ckey) if ckey else None))
            t=tot.setdefault(MM,[0,0,0,0]); t[0]+=sp; t[1]+=num(g('orders')); t[2]+=num(g('buy')); t[3]+=num(g('rev'))
    return tot

el=scan(EL,[('Август','2026-08'),('Сентябрь','2026-09')],'Елена')
ma=scan(MAR,[('Июль','2026-07'),('Август','2026-08'),('Сентябрь','2026-09')],'Мария')

# ---------- АНЯ: внешний трафик (июнь–июль, Wildberries) ----------
wn=openpyxl.load_workbook(ANYA,read_only=True,data_only=True)
def pdate(x):
    if isinstance(x,datetime.datetime): return x
    m=re.match(r'^\s*(\d{1,2})[./](\d{1,2})',txt(x))
    try: return datetime.datetime(2026,int(m.group(2)),int(m.group(1))) if m else None
    except: return None
for sh,prod,ci,oi in [('Отчет Триггер','atf',18,21),('Отчет аппликатор','appl',18,22)]:
    ws=wn[sh]
    for r in ws.iter_rows(min_row=2,values_only=True):
        if not r or len(r)<=oi: continue
        nm=txt(r[4])
        if not nm or nm.startswith(('Блогер','ИТОГ','Итог')): continue
        d=pdate(r[2])
        if not d or d.month not in (6,7): continue
        sp=num(r[9]); o=num(r[oi]); cl=num(r[ci])
        if sp==0 and o==0 and cl==0: continue
        w=blank(); w[widx(d)]=[sp,o,0,0,cl,1,0,0]
        R.append(dict(m=d.strftime('%Y-%m'),prod=prod,plat='Wildberries',chan=chan(r[3]),act='Инфлюенс-посев',name=nm[:46],
            link=txt(r[15]) or txt(r[5]) or None,date=d.strftime('%d.%m'),tag='mobz',mgr='Аня',w=w,est=True))

# ---------- БОТЫ: реклама на регистрации ----------
# «Зелёный доктор» — лист «Отчет Бот» у Ани (июнь–июль, весь Telegram)
wsb=wn['Отчет Бот']; curm=None
MB={'МАРТ':3,'АПРЕЛЬ':4,'МАЙ':5,'ИЮНЬ':6,'ИЮЛЬ':7,'АВГУСТ':8,'СЕНТЯБРЬ':9}
for r in wsb.iter_rows(min_row=1,values_only=True):
    if not r or r[0] is None: continue
    a=txt(r[0]).upper()
    if a in MB: curm=MB[a]; continue
    if a=='ДОГОВОР' or curm is None: continue
    if a.isdigit(): continue          # итоговая строка блока
    if not txt(r[12]).lower().startswith('опублик'): continue
    d=pdate(r[1]) or datetime.datetime(2026,curm,15)
    if d.month!=curm: d=datetime.datetime(2026,curm,min(d.day,28))
    sp=num(r[8]); rg=num(r[16]); rch=num(r[13])
    if sp==0 and rg==0: continue
    w=blank(); w[widx(d)]=[sp,0,0,0,0,1,0,rg]
    R.append(dict(m=d.strftime('%Y-%m'),prod='gd',plat='Бот',chan='Telegram',act='Инфлюенс-посев',
        name=txt(r[3]).replace('\n',' ').strip()[:46] or 'без названия',
        link=txt(r[4]) or None,date=d.strftime('%d.%m'),tag='botlink',mgr='Аня',w=w,est=False,reach=rch))

# «Умный сад» — отдельный файл закупки (март–май, весь Telegram, метка startapp=src_*)
import csv as _csv
US=U+'06_bot_umny_sad.csv'
try:
    rows=list(_csv.reader(open(US,encoding='utf-8')))
    for r in rows[3:]:
        if len(r)<13 or not txt(r[0]).lower().startswith('реклама заверш'): continue
        d=pdate(r[5].strip().replace(',','.'))
        if not d: continue
        sp=num(r[3]); rg=num(r[12])
        if sp==0 and rg==0: continue
        w=blank(); w[widx(d)]=[sp,0,0,0,num(r[9]),1,0,rg]
        nm=txt(r[1]).replace('https://t.me/','@').replace('@+','приват ')
        R.append(dict(m=d.strftime('%Y-%m'),prod='us',plat='Бот',chan='Telegram',act='Инфлюенс-посев',
            name=nm[:46],link=txt(r[1]) or None,date=d.strftime('%d.%m'),
            tag='botlink',mgr='Татьяна',w=w,est=False))
except Exception as e: print('умный сад skip',e)

# ---------- АНАСТАСИЯ: внутренняя реклама, самовыкупы, рейтинг ----------
wa=openpyxl.load_workbook(NAS,read_only=True,data_only=True); ws=wa['Отчёт по неделям']
MN={'ИЮЛЬ':'2026-07','АВГУСТ':'2026-08','СЕНТЯБРЬ':'2026-09'}
acc={}; selfb={}; allord={}; RATE={}
cur=None; wk=None
for r in ws.iter_rows(min_row=24,values_only=True):
    if not r or not r[0]: continue
    a=txt(r[0])
    if a.startswith('▼'):
        cur=MN.get(a.split()[1] if len(a.split())>1 else '',None); wk=None; continue
    m=re.match(r'Неделя (\d)',a)
    if m:
        wk=int(m.group(1))-1
        if cur:
            allord.setdefault(cur,[0]*5)[wk]+=num(r[4]); selfb.setdefault((cur,'all'),[0]*5)[wk]+=num(r[5])
        continue
    if not cur or wk is None: continue
    k=a.upper()
    if not k.startswith(('CPC','CPM','ОЗОН')): continue
    pl='Ozon' if k.startswith('ОЗОН') else 'Wildberries'
    prod=prodof(r[1])
    key=(cur,pl,prod); acc.setdefault(key,blank())
    acc[key][wk]=[acc[key][wk][0]+num(r[2]),acc[key][wk][1]+num(r[6]),acc[key][wk][2]+num(r[9]),
                  acc[key][wk][3]+num(r[11]),0,0,0,0]
    selfb.setdefault((cur,pl),[0]*5)[wk]+=num(r[5])
    if len(r)>19 and r[19] not in (None,''):
        RATE.setdefault(cur,{}).setdefault(pl,{}).setdefault(prod,[None]*5)[wk]=round(num(r[19]),1)
    if pl=='Ozon': allord.setdefault(cur+'|oz',[0]*5)[wk]+=num(r[4])
for (mm,pl,prod),w in acc.items():
    if sum(x[0]+x[1] for x in w)==0: continue
    R.append(dict(m=mm,prod=prod,plat=pl,chan='Внутренняя реклама '+('WB' if pl=='Wildberries' else 'Ozon'),act='Внутренняя реклама',
        sys=True,name='Кампании CPC + CPM' if pl=='Wildberries' else 'Продвижение в поиске',
        link=None,date='весь месяц',tag='sys',mgr='Анастасия',w=w,est=False))
for (mm,pl),arr in selfb.items():
    if pl=='all' or sum(arr)==0: continue
    w=blank()
    for i,v in enumerate(arr): w[i]=[0,v,v,0,0,0,0,0]
    R.append(dict(m=mm,prod='atf',plat=pl,chan='Самовыкупы',act='Самовыкупы',sys=True,info=True,
        name='Самовыкупы менеджера по маркетплейсам — в итог не входят',link=None,date='весь месяц',
        tag='sys',mgr='Анастасия',w=w,est=False))

# 1) внутренняя реклама: расход и заказы из API, разложены по недельной сетке Анастасии
def scale(rows,newtot,idx):
    cur=sum(r['w'][i][idx] for r in rows for i in range(5))
    if not cur:
        if rows: rows[0]['w'][0][idx]=newtot
        return
    kf=newtot/cur; tot=0; last=None
    for r in rows:
        for i in range(5):
            v=r['w'][i][idx]
            if v:
                nv=round(v*kf); r['w'][i][idx]=nv; tot+=nv; last=(r,i)
    if last and tot!=newtot: last[0]['w'][last[1]][idx]+=newtot-tot
for mm,src,pl in [(m,A_WBADS,'Wildberries') for m in A_WBADS]+[(m,A_OZADS,'Ozon') for m in A_OZADS]:
    d=src.get(mm) or {}
    if not d or mm not in MONTHS: continue
    rows=[r for r in R if r['m']==mm and r['plat']==pl and r.get('act')=='Внутренняя реклама']
    if not rows:   # в отчёте Анастасии этого месяца нет — заводим строку по API
        w=blank(); w[0]=[0,0,0,0,0,0,0,0]
        rows=[dict(m=mm,prod='atf',plat=pl,chan='Внутренняя реклама '+('WB' if pl=='Wildberries' else 'Ozon'),
             act='Внутренняя реклама',sys=True,name='Кампании площадки — данные из API',link=None,
             date='весь месяц',tag='sys',mgr='Анастасия',w=w,est=False)]
        R.extend(rows)
    if d.get('spend'): scale(rows,d['spend'],0)
    if d.get('orders'): scale(rows,d['orders'],1)
    if d.get('sum'):    scale(rows,d['sum'],6)
    for r in rows: r['src']='api'

# 3) Ozon: заказы и сумма заказов по размещениям — из отчёта по внешнему трафику.
#    Ключ ищем в трёх разрезах Ozon: campaign, content, term — метки у менеджеров лежат по-разному.
SRCMAP={'camp':A_UTM,'cont':(A_DEEP.get('content') or {}),'term':(A_DEEP.get('term') or {})}
UTMHIT={}
allkeys=set()
for r in R:
    if r.get('ckey'): allkeys.add((r['m'],tuple(r['ckey'])))
for mm,(kind,val) in sorted(allkeys):
    d=(SRCMAP.get(kind) or {}).get(mm,{}).get(val)
    if not d: continue
    rows=[r for r in R if r['m']==mm and r.get('ckey')==[kind,val] and not r.get('sys')]
    if not rows: continue
    UTMHIT[(mm,kind,val)]=[r['name'] for r in rows]
    for idx,vv in ((1,d.get('o_attr',0)),(6,d.get('sum',0))):
        base=[sum(r['w'][i][1] for i in range(5)) or 0 for r in rows]
        tb=sum(base); run=0
        for j,r in enumerate(rows):
            share=(base[j]/tb) if tb else 1.0/len(rows)
            v=round(vv*share) if j<len(rows)-1 else round(vv-run)
            run+=v
            wi=next((i for i in range(5) if r['w'][i][0] or r['w'][i][5]),0)
            for i in range(5): r['w'][i][idx]=0
            r['w'][wi][idx]=max(0,v)
    for r in rows: r['src']='api'; r['est']=False

# ---------- ФАКТ ПЛОЩАДОК из выгрузок WB и Ozon ----------
FACT={}   # (мес, площадка, продукт) -> [[зак,вык,сумма] x5]
def fblank(): return [[0,0,0] for _ in range(5)]
wbf=openpyxl.load_workbook(U+'07_wb_export.xlsx',data_only=True)['Товары']
for i in range(3,wbf.max_row+1):
    art=txt(wbf.cell(i,1).value); d=wbf.cell(i,10).value
    if not d: continue
    d=d if isinstance(d,datetime.datetime) else datetime.datetime.strptime(str(d)[:10],'%Y-%m-%d')
    mm=d.strftime('%Y-%m')
    if mm not in MONTHS: continue
    prod='appl' if 'applik' in art.lower() else 'atf'
    k=(mm,'Wildberries',prod); FACT.setdefault(k,fblank())
    w=FACT[k][widx(d)]
    w[0]+=num(wbf.cell(i,16).value); w[1]+=num(wbf.cell(i,17).value); w[2]+=num(wbf.cell(i,22).value)
ozf=openpyxl.load_workbook(U+'08_ozon_export.xlsx',data_only=True)['По товарам']
for i in range(14,ozf.max_row+1):
    d=ozf.cell(i,10).value
    if not d: continue
    d=d if isinstance(d,datetime.datetime) else datetime.datetime.strptime(str(d)[:10],'%Y-%m-%d')
    mm=d.strftime('%Y-%m')
    if mm not in MONTHS: continue
    k=(mm,'Ozon','atf'); FACT.setdefault(k,fblank())
    w=FACT[k][widx(d)]
    w[0]+=num(ozf.cell(i,15).value); w[1]+=num(ozf.cell(i,16).value); w[2]+=num(ozf.cell(i,13).value)

# Ozon: свежая сводка за 01–07.09 (в дневной выгрузке было только по 04.09)
try:
    oz3=openpyxl.load_workbook(U+'09_ozon_export_fresh.xlsx',data_only=True)['По товарам']
    new=[num(oz3.cell(13,16).value),num(oz3.cell(13,18).value),num(oz3.cell(13,12).value)]
    k=('2026-09','Ozon','atf')
    if k in FACT and sum(new)>0:
        cur=[sum(x[j] for x in FACT[k]) for j in range(3)]
        d=[new[j]-cur[j] for j in range(3)]
        if d[0]>0:   # дни 5–7 сентября: 5 и 6 в неделю 1, 7 — в неделю 2
            for j in range(3):
                FACT[k][0][j]+=round(d[j]*2/3); FACT[k][1][j]+=round(d[j]/3)
        print('Ozon сентябрь дополнен до',new,'(было',cur,')')
except Exception as e: print('oz3 skip',e)

# 2) итоги Wildberries — из воронки WB API, недельная разбивка сохраняется пропорционально
BUYSUM={}
for mm,per in A_WBFUN.items():
    for prod,d in per.items():
        k=(mm,'Wildberries',prod)
        if k not in FACT: FACT[k]=fblank()
        F=FACT[k]
        for j,key in ((0,'orders'),(1,'buyouts'),(2,'sum')):
            newtot=d.get(key)
            if newtot is None: continue
            cur=sum(x[j] for x in F)
            if cur:
                kk=newtot/cur; tot2=0; li=0
                for i in range(5):
                    if F[i][j]:
                        F[i][j]=round(F[i][j]*kk); tot2+=F[i][j]; li=i
                if tot2!=newtot: F[li][j]+=newtot-tot2
            else:
                F[0][j]=newtot
        if d.get('buysum'): BUYSUM[k]=d['buysum']

# 4) итоги Ozon — из API: количество товаров и сумма ДО скидки площадки (без СПП)
for mm,d in A_OZORD.items():
    k=(mm,'Ozon','atf')
    if k not in FACT: continue
    F=FACT[k]
    qty=sum(v.get('qty',0) for v in (d.get('items') or {}).values()) or d.get('ship')
    for j,newtot in ((0,qty),(2,d.get('sum'))):
        if not newtot: continue
        cur=sum(x[j] for x in F)
        if cur:
            kk=newtot/cur; run=0; li=0
            for i in range(5):
                if F[i][j]:
                    F[i][j]=round(F[i][j]*kk); run+=F[i][j]; li=i
            if run!=newtot: F[li][j]+=newtot-run
        else: F[0][j]=newtot

AVG={k:(sum(x[2] for x in v)/sum(x[0] for x in v) if sum(x[0] for x in v) else 0) for k,v in FACT.items()}
AVGBUY={k:(BUYSUM[k]/sum(x[1] for x in FACT[k]) if sum(x[1] for x in FACT[k]) else 0) for k in BUYSUM}
for r in R:
    if not r.get('sys'): continue
    a=AVG.get((r['m'],r['plat'],r['prod']),0)
    if not a: continue
    for i in range(5): r['w'][i][6]=round(r['w'][i][1]*a)

# ---------- органика = факт − реклама − посевы − самовыкупы (по неделям) ----------
ORG={}
for (mm,pl,prod),F in sorted(FACT.items()):
    tot_o=sum(x[0] for x in F); tot_b=sum(x[1] for x in F); tot_s=sum(x[2] for x in F)
    ad_s=sum(sum(a[i][6] for i in range(5)) for (m2,p2,pr),a in acc.items() if (m2,p2,pr)==(mm,pl,prod))
    sd_s=sum(sum(x[6] for x in r['w']) for r in R if r['m']==mm and r['plat']==pl and r['prod']==prod and not r.get('sys'))
    sb_s=round(sum(selfb.get((mm,pl),[0]*5))*(tot_s/tot_o if tot_o else 0)) if prod=='atf' else 0
    org_s=max(0,tot_s-ad_s-sd_s-sb_s)
    avg = AVGBUY.get((mm,pl,prod)) or (tot_s/tot_o if tot_o else 0)
    ad_o=sum(sum(a[i][1] for i in range(5)) for (m2,p2,pr),a in acc.items() if (m2,p2,pr)==(mm,pl,prod))
    ad_b=sum(sum(a[i][2] for i in range(5)) for (m2,p2,pr),a in acc.items() if (m2,p2,pr)==(mm,pl,prod))
    sd_o=sum(sum(x[1] for x in r['w']) for r in R if r['m']==mm and r['plat']==pl and r['prod']==prod and not r.get('sys'))
    sd_b=sum(sum(x[2] for x in r['w']) for r in R if r['m']==mm and r['plat']==pl and r['prod']==prod and not r.get('sys'))
    sb=sum(selfb.get((mm,pl),[0]*5)) if prod=='atf' else 0
    o=max(0,tot_o-ad_o-sd_o-sb); bq=max(0,tot_b-ad_b-sd_b-sb)
    if o==0 and bq==0: continue
    base=[x[0] for x in F]; tb=sum(base) or 1
    w=blank(); ao=ab=asu=0
    for i in range(5):
        oi=round(o*base[i]/tb) if i<4 else round(o-ao)
        bi=round(bq*base[i]/tb) if i<4 else round(bq-ab)
        si=round(org_s*base[i]/tb) if i<4 else round(org_s-asu)
        ao+=oi; ab+=bi; asu+=si
        w[i]=[0,max(0,oi),max(0,bi),round(max(0,bi)*avg),0,0,max(0,si),0]
    ORG[(mm,pl,prod)]=[o,bq]
    R.append(dict(m=mm,prod=prod,plat=pl,chan='Органика',act='Органика',sys=True,
        name='Остаток: факт площадки минус внутренняя реклама, посевы и самовыкупы',
        link=None,date='весь месяц',tag='sys',mgr='—',w=w,est=False))
for mm in MONTHS:
    if any(r['m']==mm and r['plat']=='Сайт trigger-fish.ru' for r in R):
        R.append(dict(m=mm,prod='atf',plat='Сайт trigger-fish.ru',chan='Органика',act='Органика',sys=True,
            name='Нет данных — заказы Тильды не выгружаются',link=None,date='—',tag='sys',mgr='—',w=blank(),est=True))

# ---------- ПЛАН: только из общего медиаплана, лист «План упрощённый» ----------
MP=U+'05_mediaplan.xlsx'
PRODKEY={'АМИНА ТРИГГЕР ФИШ':'atf','АППЛИКАТОР ЯРМАКОВА':'appl',
         'БОТ «ЗЕЛЁНЫЙ ДОКТОР»':'gd','БОТ «УМНЫЙ САД»':'us','БОТ «КЛЁВОЕ МЕСТО»':'km'}
MONCOL={2:'2026-07',3:'2026-08',4:'2026-09',5:'2026-10',6:'2026-11',7:'2026-12'}
SKIP={'DRR, %','CPO средний, ₽','Средняя цена регистрации, ₽','+ свободный канал','-','Итого бюджет, ₽'}
ACT2CHAN={'ВБ внутренняя':('chan','Внутренняя реклама WB'),'Ozon внутренняя':('chan','Внутренняя реклама Ozon'),
          'Авито':('plat','Авито')}
PLANS={}
try:
    mws=openpyxl.load_workbook(MP,data_only=True)['План упрощённый']
    mrows=[list(r) for r in mws.iter_rows(min_row=1,values_only=True)]
    cur=None; sec=None
    for r in mrows:
        a=txt(r[0])
        if a.upper().startswith('СВОД ПО КАНАЛАМ'): break
        if a.startswith('▼'):
            t=a[1:].strip().upper(); cur=None
            for k,v in PRODKEY.items():
                if t.startswith(k): cur=v
            sec=None; continue
        if a.startswith('──'): sec=a; continue
        if not cur: continue
        for ci,mm in MONCOL.items():
            P=PLANS.setdefault(mm,{'prod':{},'plat':{},'chan':{},'act':{}})
            d=P['prod'].setdefault(cur,{})
            v=num(r[ci]) if ci<len(r) else 0
            if a=='Итого бюджет, ₽': d['b']=round(v)
            elif a=='Итого заказов': d['o']=round(v)
            elif a=='Итого регистраций': d['r']=round(v)
            elif a=='Выкупы': d['buy']=round(v)
            elif a=='Выручка план, ₽': d['rev']=round(v)
            elif sec and 'Бюджет по каналам' in sec and a and a not in SKIP:
                if v: P['act'].setdefault(a,{}).setdefault(cur,{})['b']=round(v)
            elif sec and ('Заказы по каналам' in sec or 'Регистрации (расч' in sec) and a and a not in SKIP:
                if v: P['act'].setdefault(a,{}).setdefault(cur,{})['o']=round(v)
    # активности -> план по площадкам и каналам дашборда
    for mm,P in PLANS.items():
        for act,per in P['act'].items():
            kind,name=ACT2CHAN.get(act,(None,None))
            if act=='Инфлюенс-маркетинг' or not kind: continue
            for prod,vv in per.items():
                if kind=='chan': P['chan'].setdefault(prod,{})[name]=dict(vv)
                elif kind=='plat': P['plat'].setdefault(prod,{})[name]=dict(vv)
                else: P['chan'].setdefault(prod,{})[name]=dict(vv)
    for mm in list(PLANS):
        P=PLANS[mm]
        P['prod']={k:v for k,v in P['prod'].items() if v.get('b') or v.get('o') or v.get('r')}
        P['act']={k:v for k,v in P['act'].items() if v}
except Exception as e:
    print('план не разобран:',e); PLANS={}

# ---------- Авито: площадка продаж с внутренней рекламой, факта пока нет ----------
ACTMAP={'Инфлюенс-маркетинг':'Инфлюенс-посев','ВБ внутренняя':'Внутренняя реклама','Ozon внутренняя':'Внутренняя реклама',
        'Авито':'Авито','SEO':'SEO-продвижение','Конкурс репостов в вк':'Конкурс репостов в ВК',
        'YouTube-канал':'Свой YouTube-канал','Бартерная реклама':'Бартер','Таргет ВК':'Таргет ВК',
        'Контекст (Директ)':'Контекст'}
seen=set()
for mm,P in PLANS.items():
    if mm not in MONTHS: continue
    for prod,pp in P['plat'].items():
        if 'Авито' in pp and (mm,prod) not in seen:
            seen.add((mm,prod))
            R.append(dict(m=mm,prod=prod,plat='Авито',chan='Внутренняя реклама Авито',act='Авито',sys=True,
                name='План есть, фактических данных пока нет',link=None,date='—',
                tag='sys',mgr='—',w=blank(),est=True))

# 5) рейтинг карточек — из API, ставится точкой на текущую неделю поверх ручных значений
if A_RATE.get('month') in MONTHS:
    _mm=A_RATE['month']; _wi=int(A_RATE.get('week_index',0))
    for _pl in ('Wildberries','Ozon'):
        for _prod,_v in (A_RATE.get(_pl) or {}).items():
            if not _v: continue
            RATE.setdefault(_mm,{}).setdefault(_pl,{}).setdefault(_prod,[None]*5)[_wi]=round(float(_v),1)
    RATESRC={'date':A_RATE.get('date'),'month':_mm,'week':_wi}
else:
    RATESRC=None

# ---------- KPI инфлюенс-менеджеров ----------
KPIFILES=[(EL,'Елена','2026-08',[('Август',8),('Сентябрь',9)]),
          (MAR,'Мария','2026-07',[('Июль',7),('Август',8),('Сентябрь',9)])]
KPICHK=[('spend','итого к оплате'),('reach','охват'),('clm','клики мобзио'),('clw','клики wb'),
        ('o','заказы'),('buy','выкупы'),('rev','выручка от выкупов'),('osum','сумма заказов')]
def kday(v):
    if isinstance(v,datetime.datetime): return v.day
    m=re.match(r'(\d{1,2})[/.](\d{1,2})',txt(v))
    return int(m.group(1)) if m else None
def lvl_cpc(c):
    if c is None: return None,None
    if c<80: return 4,1.0
    if c<=100: return 3,0.7
    if c<=149: return 2,0.5
    return 1,0.2
def lvl_cpv(v):
    if v is None: return None,None
    if v<=0.49: return 4,1.0
    if v<=1: return 3,0.7
    if v<=1.3: return 2,0.5
    return 1,0.1
KPI={}
for path,who,first,shs in KPIFILES:
    wbk=openpyxl.load_workbook(path,read_only=True,data_only=True)
    pubs=[]; formerr={}
    for sh,mon in shs:
        rws=[list(r) for r in wbk[sh].iter_rows(min_row=5,values_only=True)]
        hdr=rws[0]; C=cmap(hdr); ci={}
        for i,c in enumerate(hdr):
            t=txt(c).lower().replace('\n',' ')
            for k,p2 in KPICHK:
                if t.startswith(p2) and k not in ci: ci[k]=i
        mm='2026-%02d'%mon
        det=[]; allr=[]; ecount=0; elist=[]
        for r in rws[1:]:
            a=txt(r[0] if r else '')
            if not a: continue
            up=a.upper()
            if up.startswith('НЕДЕЛЯ'): det=[]; continue
            if up.startswith('ИТОГО') or a.lower().startswith('итого'):
                isM='ЗА МЕСЯЦ' in up; src=allr if isM else det
                for k,i in ci.items():
                    got=num(r[i] if i<len(r) else None)
                    exp=sum(num(x[i] if i<len(x) else None) for x in src)
                    if exp==0 and got==0: continue
                    if abs(got-exp)>max(1,0.01*max(abs(exp),abs(got))):
                        ecount+=1
                        elist.append(a[:30]+' · '+k+': должно '+format(round(exp),',').replace(',',' ')+', стоит '+format(round(got),',').replace(',',' '))
                if not isM: det=[]
                continue
            det.append(r); allr.append(r)
            g=lambda k: r[C[k]] if k in C and C[k]<len(r) else None
            if not done(g('status')): continue
            pubs.append(dict(mon=mon,day=kday(g('date')),nm=txt(r[0])[:38],
                sp=num(g('spend')),clm=num(g('clm')),rch=num(g('reach'))))
        formerr[mm]={'n':ecount,'list':elist[:8]}
    months={}
    for _,mon in shs:
        mm='2026-%02d'%mon
        if mm==first:
            win=[p for p in pubs if p['mon']==mon and p['day'] and p['day']<=17]
            wlab='1–17 '+MONTHS[mm]['label'].split()[0].lower()
        else:
            pmon=mon-1
            win=[p for p in pubs if (p['mon']==pmon and p['day'] and p['day']>=18) or (p['mon']==mon and p['day'] and p['day']<=17)]
            pl=MONTHS.get('2026-%02d'%pmon,{}).get('label','').split()[0].lower()
            wlab='18–31 '+pl+' + 1–17 '+MONTHS[mm]['label'].split()[0].lower()
        lk=[p for p in win if p['clm']>0]
        cl=sum(p['clm'] for p in lk); spl=sum(p['sp'] for p in lk)
        cpc=(spl/cl) if cl else None
        rc=[p for p in win if p['rch']>0]
        rch=sum(p['rch'] for p in rc); spr=sum(p['sp'] for p in rc)
        cpv=(spr/rch) if rch else None
        L2,K2=lvl_cpc(cpc); L3,K3=lvl_cpv(cpv)
        months[mm]=dict(win=wlab,pubs=len(win),spend=round(sum(p['sp'] for p in win)),
            cpc=(round(cpc,1) if cpc else None),cpcRows=len(lk),cpcClicks=round(cl),cpcSpend=round(spl),
            cpv=(round(cpv,3) if cpv else None),cpvRows=len(rc),cpvReach=round(rch),cpvSpend=round(spr),
            L2=L2,K2=K2,L3=L3,K3=K3,err=formerr.get(mm,{'n':0,'list':[]}),
            rows=sorted([dict(nm=p['nm'],d=('%02d.%02d'%(p['day'],p['mon'])) if p['day'] else '—',
                sp=round(p['sp']),clm=round(p['clm']),rch=round(p['rch'])) for p in win],
                key=lambda x:-x['sp'])[:40])
    KPI[who]=dict(first=first,months=months)

# последний рейтинг
last={}
for mm in sorted(RATE):
    for pl,pr in RATE[mm].items():
        for prod,arr in pr.items():
            for i,v in enumerate(arr):
                if v: last[pl]={'v':v,'when':MONTHS[mm]['weeks'][i],'prod':prod}
meta=dict(months=MONTHS,kpi=KPI,ratingApi=RATESRC,ratingContent=(A_RATE.get('content') or {}),plans=PLANS,rating=RATE,ratingLast=last,
  pend={f'{k[0]}|{k[1]}':v for k,v in PEND.items()},
  selfbuy={'wb':round(sum(selfb.get(('2026-08','Wildberries'),[0]*5))),'oz':round(sum(selfb.get(('2026-08','Ozon'),[0]*5)))},
  fact={f'{k[0]}|{k[1]}|{k[2]}':[round(sum(x[0] for x in v)),round(sum(x[1] for x in v)),round(sum(x[2] for x in v))] for k,v in FACT.items()},
  actmap=ACTMAP,
  ozutm={mm:sorted([dict(camp=c,**v,
        who=[r['name'] for r in R if r.get('camp')==c and r['m']==mm])
      for c,v in per.items()],key=lambda x:-x['o_attr'])
    for mm,per in A_UTM.items()},
  utmhit={f'{k[0]}|{k[1]}|{k[2]}':v for k,v in UTMHIT.items()},
  apisrc=dict(wb_funnel=sorted(A_WBFUN),wb_ads=sorted(A_WBADS),oz_ads=sorted(A_OZADS),oz_utm=sorted(A_UTM)),
  check=dict(el=el.get('2026-08'),ma=ma.get('2026-08')))
DATA=json.dumps(dict(records=R,meta=meta),ensure_ascii=False)
json.dump(dict(records=R,meta=meta),open('records.json','w'),ensure_ascii=False)
tpl=open('template.html',encoding='utf-8').read()
open('../dashboard-influence.html','w',encoding='utf-8').write(tpl.replace('__DATA__',DATA))
print('dashboard-influence.html собран')
from collections import Counter
print('записей',len(R),'| месяцы',Counter(r['m'] for r in R))
print('Елена',{k:[round(x) for x in v] for k,v in el.items()})
print('Мария',{k:[round(x) for x in v] for k,v in ma.items()})
print('ожидают',PEND)
print('рейтинг последний',last)
print('ФАКТ площадок (зак/вык/сумма):')
for k in sorted(FACT): print('  ',k,[round(sum(x[j] for x in FACT[k])) for j in range(3)])
print('ОРГАНИКА (зак/вык):')
for k in sorted(ORG): print('  ',k,[round(x) for x in ORG[k][:2]])
