# -*- coding: utf-8 -*-
"""
Сборщик дашборда инфлюенс-маркетинга. Пересобран 02.10.2026 после утраты build3.py.

Как устроен:
  • закрытые месяцы (март–август) берутся готовыми из frozen.json и НЕ пересчитываются;
  • текущие месяцы (сентябрь, октябрь) считаются из таблиц менеджеров, медиаплана и API;
  • на выходе records.json, который mk.py подставляет в tpl.html.

Запуск:  python3 build.py
"""
import openpyxl, json, re, datetime, os, csv as _csv
from collections import Counter as _Cnt

TODAY = datetime.datetime(2026, 10, 2)
LIVE  = ['2026-09', '2026-10']          # месяцы, которые считаем
SRC   = 'src/'
API   = 'api/'

MONTHS={'2026-03':{'label':'Март 2026','weeks':['1–7 мар','8–14 мар','15–21 мар','22–28 мар','29–31 мар']},
        '2026-04':{'label':'Апрель 2026','weeks':['1–7 апр','8–14 апр','15–21 апр','22–28 апр','29–30 апр']},
        '2026-05':{'label':'Май 2026','weeks':['1–7 май','8–14 май','15–21 май','22–28 май','29–31 май']},
        '2026-06':{'label':'Июнь 2026','weeks':['1–7 июн','8–14 июн','15–21 июн','22–28 июн','29–30 июн']},
        '2026-07':{'label':'Июль 2026','weeks':['1–7 июл','8–14 июл','15–21 июл','22–28 июл','29–31 июл']},
        '2026-08':{'label':'Август 2026','weeks':['1–9 авг','10–16 авг','17–23 авг','24–30 авг','31 авг']},
        '2026-09':{'label':'Сентябрь 2026','weeks':['1–6 сен','7–13 сен','14–20 сен','21–27 сен','28–30 сен']},
        '2026-10':{'label':'Октябрь 2026','weeks':['1–7 окт','8–14 окт','15–21 окт','22–28 окт','29–31 окт']}}

def widx(d):
    dd=d.day
    if d.month==8: return 0 if dd<=9 else 1 if dd<=16 else 2 if dd<=23 else 3 if dd<=30 else 4
    if d.month==9: return 0 if dd<=6 else 1 if dd<=13 else 2 if dd<=20 else 3 if dd<=27 else 4
    return 0 if dd<=7 else 1 if dd<=14 else 2 if dd<=21 else 3 if dd<=28 else 4

def num(x):
    if x is None: return 0
    if isinstance(x,(int,float)): return float(x)
    s=re.sub(r'[^0-9.\-]','',str(x).replace('\xa0','').replace(' ','').replace(',','.'))
    try: return float(s) if s not in('','-','.') else 0
    except: return 0
def txt(x): return ('' if x is None else str(x)).strip()
def blank(): return [[0,0,0,0,0,0,0,0] for _ in range(5)]
# позиции в неделе: 0 расход, 1 заказы, 2 выкупы, 3 выручка, 4 клики, 5 размещений, 6 сумма заказов, 7 регистрации

CH={'max':'MAX','макс':'MAX','vk':'VK','vkontakte':'VK','вк':'VK','inst':'Instagram','instagram':'Instagram',
    'youtube':'YouTube','telegram':'Telegram','tiktok':'TikTok','rutube':'Rutube','тг':'Telegram',
    'телеграм':'Telegram','телега':'Telegram','ютуб':'YouTube','тик-ток':'TikTok','vk клипы':'VK','vk видео':'VK'}
def chan(s):
    k=txt(s).lower().strip()
    if k in ('wb','вб','ozon','озон','сайт',''): return 'Канал не указан'
    return CH.get(k, txt(s) or 'Канал не указан')
def plat(s):
    k=txt(s).lower()
    if 'wb' in k or 'вб' in k: return 'Wildberries'
    if 'ozon' in k or 'озон' in k: return 'Ozon'
    if 'сайт' in k: return 'Сайт trigger-fish.ru'
    return 'Не указано'
def platByTag(art,utm,link,promo):
    a=txt(art).upper(); u=(txt(utm)+' '+txt(link)).lower()
    if a.startswith('WW') or 'wildberries' in u: return 'Wildberries'
    if 'ozon' in u: return 'Ozon'
    if 'trigger-fish' in u: return 'Сайт trigger-fish.ru'
    return 'Не указано'
def tagfor(pl,art,utm,promo):
    art,utm,promo=txt(art),txt(utm),txt(promo)
    if promo in ('нет','-','—','-TRG'): promo=''
    if art in ('нет','-','—'): art=''
    if pl.startswith('Сайт'): return 'utm' if utm else ('promo' if promo else 'none')
    if pl=='Wildberries': return 'art' if art.upper().startswith('WW') else ('promo' if promo else 'none')
    if pl=='Ozon': return 'utm' if ('utm_' in utm or 'onelink.me' in utm) else ('promo' if promo else 'none')
    return 'none'

BOTS={'gd':'Бот «Зелёный доктор»','us':'Бот «Умный сад»','km':'Бот «Клёвая карта»'}
PRODS={'atf':'AMINOTRIGGERFISH Сигнал','atfp':'AMINOTRIGGERFISH Прикормка','appl':'Аппликатор Ермакова'}
def prodof(x):
    k=txt(x).lower()
    if 'зелён' in k or 'зелен' in k or 'доктор' in k or 'green_doctor' in k: return 'gd'
    if 'умный сад' in k or 'умная гряд' in k or 'garden_app' in k: return 'us'
    if 'прикорм' in k: return 'atfp'
    if 'клёв' in k or 'клев' in k: return 'km'
    if 'апплик' in k: return 'appl'
    return 'atf'
def done(x): return txt(x).lower().startswith(('вышл','опублик'))

PATS=[('spend','итого к оплате'),('osum','сумма заказов'),('orders','заказы'),('buy','выкупы'),
      ('rev','выручка от выкупов'),('clm','клики мобзио'),('clw','клики wb'),('reach','охват'),
      ('cart','корзины'),('date','дата выхода'),('status','статус'),('plat','площадка'),
      ('sales','канал продаж'),('art','подменный артикул'),('utm','ссылка с utm'),('promo','промокод'),
      ('link','ссылка на размещение'),('prod','продукт'),('scn','сценарий'),('reg','регистрации')]
def cmap(hdr):
    m={}
    for i,c in enumerate(hdr):
        t=txt(c).lower().replace('\n',' ')
        for k,p in PATS:
            if t.startswith(p) and k not in m: m[k]=i
    return m

def loadapi(n):
    try: return json.load(open(API+n))
    except Exception as e:
        print('  api/%s не прочитан: %s'%(n,e)); return {}
A_UTM   = loadapi('ozon_utm.json')
A_DEEP  = loadapi('ozon_utm_deep.json')
A_OZADS = loadapi('ozon_ads.json')
A_OZORD = loadapi('ozon_orders.json')
A_WBADS = loadapi('wb_ads.json')
A_WBFUN = loadapi('wb_funnel.json')
A_AVITO = loadapi('avito.json')
A_RATE  = loadapi('rating.json')
A_SHORT = loadapi('shortlinks.json')

# ---------- РАЗБОР МЕТОК ----------
GENERIC={'vendor_org_4202727',''}
def badkey(v):
    """мусорный ключ метки: слишком короткий или из одних цифр — такому верить нельзя"""
    v=txt(v)
    return (not v) or len(v)<4 or v.isdigit() or v in GENERIC

def parts(u, own=True):
    """разбирает ссылку на части метки.
       own=False — чужая ссылка (на публикацию), из неё параметр c= брать НЕЛЬЗЯ:
       у vk.ru это собственный служебный параметр, а не наша метка."""
    u=txt(u)
    if not u: return {}
    m=re.search(r'onelink\.me/SNMZ/([A-Za-z0-9]+)',u)
    if m:
        c=A_SHORT.get(m.group(1))
        return {'camp':c,'short':m.group(1)} if c else {'short':m.group(1),'unresolved':True}
    d={}
    for k,pat in (('camp',r'utm_campaign=([^&\s]+)'),('cont',r'utm_content=([^&\s]+)'),
                  ('term',r'utm_term=([^&\s]+)'),('src',r'utm_source=([^&\s]+)'),('med',r'utm_medium=([^&\s]+)')):
        mm=re.search(pat,u)
        if mm: d[k]=mm.group(1).replace('+',' ').replace('%0A','').strip()
    if own and 'camp' not in d and 'vk.ru' not in u.lower() and 'vk.com' not in u.lower():
        mm=re.search(r'[?&]c=([^&\n]+)',u)
        if mm: d['camp']=mm.group(1).strip()
    for k in ('camp','cont','term'):
        if k in d and badkey(d[k]): del d[k]
    return d

# ---------- ТАБЛИЦЫ МЕНЕДЖЕРОВ ----------
R=[]; PEND={}; UNRESOLVED=[]
def scan(path, sheets, mgr):
    wb=openpyxl.load_workbook(path, read_only=True, data_only=True)
    tot={}
    for sh,MM in sheets:
        if sh not in wb.sheetnames:
            print('  ! лист «%s» не найден в %s'%(sh,os.path.basename(path))); continue
        ws=wb[sh]; rows=list(ws.iter_rows(min_row=5, values_only=True))
        if not rows: continue
        C=cmap(rows[0]); mon=int(MM[-2:])
        for r in rows[1:]:
            if not r or not r[0]: continue
            nm=txt(r[0])
            if nm.upper().startswith(('ИТОГО ЗА МЕСЯЦ','АВТОМАТИЧЕСКАЯ СВОДКА')): break
            if nm.startswith(('НЕДЕЛЯ','Итого','ИТОГО')): continue
            g=lambda k: r[C[k]] if k in C and C[k]<len(r) else None
            if not done(g('status')):
                p=PEND.setdefault((MM,mgr),[0,0]); p[0]+=1; p[1]+=num(g('spend')); continue
            d=g('date')
            if not isinstance(d,datetime.datetime):
                m=re.match(r'(\d{1,2})[/.](\d{1,2})',txt(d))
                d=datetime.datetime(2026,mon,int(m.group(1))) if m else datetime.datetime(2026,mon,15)
            if d.month!=mon: d=datetime.datetime(2026,mon,min(d.day,28))
            prod=prodof(g('prod'))
            pl='Бот' if prod in BOTS else plat(g('sales'))
            if pl=='Не указано': pl=platByTag(g('art'),g('utm'),g('link'),g('promo'))
            sp=num(g('spend'))
            w=blank(); w[widx(d)]=[sp,num(g('orders')),num(g('buy')),num(g('rev')),
                                  num(g('clm')) or num(g('clw')),1,num(g('osum')),num(g('reg'))]
            tg=tagfor(pl,g('art'),g('utm'),g('promo'))
            P=parts(g('utm'), own=True)
            if P.get('unresolved'):
                UNRESOLVED.append((MM,mgr,nm[:40],P['short'],txt(g('utm'))))
            if not any(k in P for k in ('camp','cont','term')):
                P2=parts(g('link'), own=False)          # запасной вариант: ссылка на публикацию
                for k in ('camp','cont','term'):
                    if k in P2: P.setdefault(k,P2[k])
            cmp_=P.get('camp'); ckey=None
            if cmp_: ckey=('camp',cmp_)
            elif P.get('cont'): ckey=('cont',P['cont'])
            elif P.get('term'): ckey=('term',P['term'])
            sc=txt(g('scn'))
            if sc.endswith('.0'): sc=sc[:-2]
            if sc in ('-','—','нет'): sc=''
            rec=dict(m=MM,prod=prod,plat=pl,chan=chan(g('plat')),act='Инфлюенс-посев',name=nm[:46],
                     link=txt(g('link')) or None,date=d.strftime('%d.%m'),tag=tg,mgr=mgr,w=w,
                     est=tg=='none',scn=sc or None,camp=cmp_,ckey=list(ckey) if ckey else None,
                     reach=num(g('reach')),subs=num(g('subs')) if 'subs' in C else 0,
                     day=d.strftime('%Y-%m-%d'),utmp={k:v for k,v in P.items() if k in ('camp','cont','term','src','med')})
            R.append(rec)
            t=tot.setdefault(MM,[0,0,0,0]); t[0]+=sp; t[1]+=num(g('orders')); t[2]+=num(g('buy')); t[3]+=num(g('rev'))
    wb.close()
    return tot

SHEETS_EL=[('Сентябрь','2026-09'),('Октябрь ','2026-10')]
SHEETS_MA=[('Сентябрь','2026-09'),('Октябрь','2026-10')]
el=scan(SRC+'elena.xlsx', SHEETS_EL, 'Елена')
ma=scan(SRC+'maria.xlsm', SHEETS_MA, 'Мария')
print('Елена :',{k:[round(x) for x in v] for k,v in el.items()})
print('Мария :',{k:[round(x) for x in v] for k,v in ma.items()})
print('ожидают выхода:',{('%s|%s'%k):[v[0],round(v[1])] for k,v in PEND.items()})
if UNRESOLVED:
    print('! НЕ РАСКРЫТЫ короткие ссылки onelink (%d шт) — метка не сопоставится:'%len(UNRESOLVED))
    for u in UNRESOLVED[:12]: print('   %s %s · %s · код %s'%(u[0],u[1],u[2],u[3]))

# ---------- АНАСТАСИЯ: берём ТОЛЬКО самовыкупы ----------
MN={'ИЮЛЬ':'2026-07','АВГУСТ':'2026-08','СЕНТЯБРЬ':'2026-09','ОКТЯБРЬ':'2026-10','НОЯБРЬ':'2026-11'}
selfb={}
try:
    nws=openpyxl.load_workbook(SRC+'anastasia.xlsx',data_only=True)['Отчёт по неделям']
    NROWS=list(nws.iter_rows(min_row=20,values_only=True))
    NC={}
    for _r in NROWS:
        if not _r or not _r[0] or not txt(_r[0]).lower().startswith('период'): continue
        for _j,_c in enumerate(_r):
            _s=txt(_c).lower().replace('ё','е').strip()
            if not _s: continue
            if 'бюджет' in _s: NC.setdefault('spend',_j)
            elif _s.startswith('охват'): NC.setdefault('reach',_j)
            elif 'всего заказов' in _s: NC.setdefault('ord',_j)
            elif 'самовыкуп' in _s: NC.setdefault('self',_j)
            elif _s.startswith('рейтинг'): NC.setdefault('rate',_j)
        break
    if 'self' not in NC:
        raise SystemExit('Таблица Анастасии: не нашёл колонку «Самовыкупы». Шапка изменилась, проверь лист «Отчёт по неделям».')
    print('колонки Анастасии:',NC)
    NG=lambda r,k:(r[NC[k]] if (k in NC and NC[k]<len(r)) else None)
    cur=None; wk=None; RATE={}
    for r in NROWS:
        if not r or not r[0]: continue
        a=txt(r[0])
        if a.startswith('▼'):
            cur=MN.get(a.split()[1] if len(a.split())>1 else '',None); wk=None; continue
        m=re.match(r'Неделя (\d)',a)
        if m:
            wk=int(m.group(1))-1
            if cur: selfb.setdefault((cur,'all'),[0]*5)[wk]+=num(NG(r,'self'))
            continue
        if not cur or wk is None: continue
        k=a.upper()
        if not k.startswith(('CPC','CPM','ОЗОН')): continue
        pl='Ozon' if k.startswith('ОЗОН') else 'Wildberries'
        selfb.setdefault((cur,pl),[0]*5)[wk]+=num(NG(r,'self'))
        _rt=NG(r,'rate')
        if _rt not in (None,''):
            RATE.setdefault(cur,{}).setdefault(pl,{}).setdefault(prodof(r[1]),[None]*5)[wk]=round(num(_rt),1)
except SystemExit: raise
except Exception as e:
    print('  ! таблица Анастасии не разобрана:',e); RATE={}
for (mm,pl),arr in list(selfb.items()):
    if pl=='all' or mm not in LIVE or sum(arr)==0: continue
    w=blank()
    for i,v in enumerate(arr): w[i]=[0,v,v,0,0,0,0,0]
    R.append(dict(m=mm,prod='atf',plat=pl,chan='Самовыкупы',act='Самовыкупы',sys=True,info=True,
        name='Самовыкупы менеджера по маркетплейсам — в итог не входят',link=None,date='весь месяц',
        tag='sys',mgr='Анастасия',w=w,est=False))
print('самовыкупы:',{('%s|%s'%k):[round(x) for x in v] for k,v in selfb.items() if k[0] in LIVE and k[1]!='all'})

# ---------- ВНУТРЕННЯЯ РЕКЛАМА: целиком из API площадок ----------
import calendar as _cal
def evenweeks(mm):
    y,m=int(mm[:4]),int(mm[5:])
    last=TODAY.day if mm==TODAY.strftime('%Y-%m') else _cal.monthrange(y,m)[1]
    w=[0.0]*5
    for day in range(1,last+1): w[widx(datetime.datetime(y,m,day))]+=1
    t=sum(w) or 1
    return [x/t for x in w]
def prodsplit(d):
    cs=d.get('campaigns') or []
    if not cs: return None
    sp={}
    for c in cs:
        n=txt(c.get('name')).lower()
        if 'applik' in n or 'апплик' in n: p='appl'
        elif 'прикорм' in n or 'prikorm' in n: p='atfp'
        else: p='atf'
        a=sp.setdefault(p,{'spend':0,'orders':0,'sum':0,'clicks':0})
        for k in a: a[k]+=num(c.get(k))
    return sp if len(sp)>1 else None

acc={}
for src,pl in ((A_WBADS,'Wildberries'),(A_OZADS,'Ozon')):
    for mm,d in src.items():
        if mm not in LIVE or not isinstance(d,dict) or not d: continue
        parts_=prodsplit(d) or {'atf':{k:num(d.get(k)) for k in ('spend','orders','sum','clicks')}}
        wk=d.get('weeks')
        for prod,v in parts_.items():
            if not (v['spend'] or v['orders']): continue
            if wk:
                tot_sp=sum(x.get('spend',0) for x in wk) or 1
                wts=[x.get('spend',0)/tot_sp for x in wk]
                tot_o=sum(x.get('orders',0) for x in wk)
                wo=[x.get('orders',0)/tot_o for x in wk] if tot_o else wts
            else:
                wts=evenweeks(mm); wo=wts
            w=blank()
            for i in range(5):
                w[i]=[round(v['spend']*wts[i]),round(v['orders']*wo[i]),0,0,
                      round(v.get('clicks',0)*wts[i]),0,round(v['sum']*wo[i]),0]
            for j,tot in ((0,v['spend']),(1,v['orders']),(6,v['sum'])):
                dlt=round(tot)-sum(x[j] for x in w)
                if dlt:
                    li=max(range(5),key=lambda i:w[i][j]); w[li][j]+=dlt
            R.append(dict(m=mm,prod=prod,plat=pl,chan='Внутренняя реклама '+('WB' if pl=='Wildberries' else 'Ozon'),
                act='Внутренняя реклама',sys=True,src='api',est=False,
                name=('Кампании WB' if pl=='Wildberries' else 'Продвижение в поиске Ozon')+' — из API',
                link=None,date='весь месяц',tag='sys',mgr='—',w=w))
            acc[(mm,pl,prod)]=w
print('внутренняя реклама:',{('%s|%s|%s'%k):round(sum(x[0] for x in v)) for k,v in acc.items()})

# ---------- ЗАКАЗЫ ПО МЕТКАМ OZON (факт площадки, не модель) ----------
# Ключ ищем в трёх разрезах отчёта: campaign, content, term — метки у менеджеров лежат по-разному.
SRCMAP={'camp':A_UTM,'cont':(A_DEEP.get('content') or {}),'term':(A_DEEP.get('term') or {})}
UTMHIT={}
def _short(v): return str(v).replace('vendor_org_4202727_','')

for r in R:
    if r.get('ckey') and r['ckey'][0]=='camp': r['camp0']=r['ckey'][1]

allkeys=set()
for r in R:
    if r.get('ckey') and r['m'] in LIVE: allkeys.add((r['m'],tuple(r['ckey'])))
for mm,(kind,val) in sorted(allkeys):
    d=(SRCMAP.get(kind) or {}).get(mm,{}).get(val)
    if not d: continue
    rows=[r for r in R if r['m']==mm and r.get('ckey')==[kind,val] and not r.get('sys') and r['plat']=='Ozon']
    if not rows: continue
    UTMHIT[(mm,kind,val)]=[r['name'] for r in rows]
    for idx,vv in ((1,d.get('o_attr',0)),(6,d.get('sum',0))):
        base=[sum(r['w'][i][1] for i in range(5)) or 0 for r in rows]
        tb=sum(base); run=0
        for j,r in enumerate(rows):
            share=(base[j]/tb) if tb else 1.0/len(rows)
            v=round(vv*share) if j<len(rows)-1 else round(vv-run)
            run+=v
            wi=next((i for i in range(5) if r['w'][i][0] or r['w'][i][5]),0)
            for i in range(5): r['w'][i][idx]=0
            r['w'][wi][idx]=max(0,v)
    for r in rows: r['src']='api'; r['est']=False

# общий пул кампании: заказы кампании, не разошедшиеся по хвостам content/term,
# делим поровну между публикациями этой кампании без своего совпадения
_grp={}
for r in R:
    if r.get('camp0') and not r.get('tail') and r['m'] in LIVE: _grp.setdefault((r['m'],r['camp0']),[]).append(r)
for (mm,cp),rows in _grp.items():
    un=[r for r in rows if (mm,*r['ckey']) not in UTMHIT]
    if not un: continue
    d=(A_UTM.get(mm) or {}).get(cp) or {}
    mo=sum(sum(w[1] for w in r['w']) for r in rows if r not in un)
    ms=sum(sum(w[6] for w in r['w']) for r in rows if r not in un)
    lo=max(0,d.get('o_attr',0)-mo); ls=max(0,d.get('sum',0)-ms)
    for idx,vv in ((1,lo),(6,ls)):
        run=0
        for j,r in enumerate(un):
            v=round(vv/len(un)) if j<len(un)-1 else round(vv-run); run+=v
            wi=next((i for i in range(5) if r['w'][i][0] or r['w'][i][5]),0)
            for i in range(5): r['w'][i][idx]=0
            r['w'][wi][idx]=max(0,v)
    for r in un:
        r['src']='api'; r['est']=False
        if len(un)>1: r['shared']=dict(key=_short(cp),n=len(un),o=lo)
    UTMHIT[(mm,'camp',cp)]=[r['name'] for r in un]
# один и тот же ключ на нескольких публикациях месяца — тоже общий пул
_same=_Cnt((r['m'],tuple(r['ckey'])) for r in R if r.get('plat')=='Ozon' and r.get('ckey')
           and not r.get('sys') and not r.get('tail') and r['m'] in LIVE)
for r in R:
    if r.get('plat')!='Ozon' or not r.get('ckey') or r.get('tail') or r.get('shared') or r['m'] not in LIVE: continue
    n=_same[(r['m'],tuple(r['ckey']))]
    if n>1:
        o=sum(sum(w[1] for w in x['w']) for x in R if x['m']==r['m'] and x.get('ckey')==r['ckey']
              and x.get('plat')=='Ozon' and not x.get('tail'))
        r['shared']=dict(key=_short(r['ckey'][1]),n=n,o=o)

# ---------- ХВОСТЫ: заказы по метке пришли в месяце, где публикации с этой меткой нет ----------
FROZ=json.load(open('frozen.json'))
ALLPUB=[r for r in R if not r.get('sys')]+[r for r in FROZ['records'] if not r.get('sys')]
TAILS=[]
for kind,S in SRCMAP.items():
    for mm,per in S.items():
        if mm not in LIVE: continue
        for val,dd in per.items():
            o=dd.get('o_attr',0)
            if not o or (mm,kind,val) in UTMHIT: continue
            prev=[r for r in ALLPUB if r.get('ckey')==[kind,val] and r['m']<mm and r.get('plat')=='Ozon']
            if not prev: continue
            src=sorted(prev,key=lambda r:(r['m'],r.get('day') or ''))[-1]
            w=blank(); w[0]=[0,o,0,0,0,0,dd.get('sum',0),0]
            t=dict(m=mm,prod=src['prod'],plat=src['plat'],chan=src['chan'],act=src['act'],
                   name=src['name']+' — заказы по метке публикации от '+src['date'],
                   link=src.get('link'),date=src['date'],tag=src['tag'],mgr=src['mgr'],w=w,
                   est=False,src='api',tail=True,camp=src.get('camp'),ckey=[kind,val],
                   camp_orig=(src['name'],src['m']))
            R.append(t); TAILS.append((mm,t['name'][:40],o))
print('хвосты прошлых месяцев:',TAILS)

# ---------- САЙТ: выгрузка корзины Тильды ----------
TEST_MAIL={'test@test.ru','test@gmail.com','vetal.wellness@gmail.com'}
TEST_PHONE='+7 (999) 999-99-99'
T_ROWS=[]
try:
    for row in _csv.DictReader(open(SRC+'tilda_orders.csv',encoding='utf-8-sig')):
        em=txt(row.get('email')).lower(); ph=txt(row.get('phone'))
        if em in TEST_MAIL or ph==TEST_PHONE: continue
        try: dt=datetime.datetime.strptime(txt(row.get('created'))[:19],'%Y-%m-%d %H:%M:%S')
        except Exception: continue
        T_ROWS.append(dict(dt=dt,em=em,ph=ph,amt=num(row.get('amount')),
                           paid=bool(txt(row.get('paymentid'))),ref=txt(row.get('referer'))))
except Exception as e:
    print('  ! Тильда не прочитана:',e)
# дубль: тот же человек, та же сумма, в пределах двух часов — остаётся оплаченный
T_ROWS.sort(key=lambda x:x['dt']); keep=[]
for x in T_ROWS:
    dup=None
    for y in keep:
        if (x['ph']and x['ph']==y['ph'] or x['em']and x['em']==y['em']) and abs(x['amt']-y['amt'])<1 \
           and abs((x['dt']-y['dt']).total_seconds())<7200: dup=y; break
    if dup is None: keep.append(x)
    elif x['paid'] and not dup['paid']: keep[keep.index(dup)]=x
TIL={}
for x in keep:
    mm=x['dt'].strftime('%Y-%m')
    if mm not in LIVE: continue
    c=re.search(r'utm_campaign=([^&\s]+)',x['ref'])
    c=c.group(1) if c else None
    if c and badkey(c): c=None
    TIL.setdefault(mm,{}).setdefault(c,[]).append(x)
for mm,per in TIL.items():
    for c,xs in per.items():
        w=blank()
        for x in xs:
            i=widx(x['dt']); w[i][1]+=1; w[i][6]+=x['amt']
            if x['paid']: w[i][2]+=1; w[i][3]+=x['amt']
        if c:
            src=re.search(r'utm_source=([^&\s]+)',xs[0]['ref'])
            R.append(dict(m=mm,prod='atf',plat='Сайт trigger-fish.ru',chan=chan(src.group(1) if src else ''),
                act='Инфлюенс-посев',name='Метка '+c+' — размещение не найдено в таблицах',
                link=None,date='—',tag='utm',mgr='—',w=w,est=False,src='tilda',camp=c))
        else:
            R.append(dict(m=mm,prod='atf',plat='Сайт trigger-fish.ru',chan='Органика',act='Органика',sys=True,
                name='Заказы с сайта без метки',link=None,date='—',tag='sys',mgr='—',w=w,est=False,src='tilda'))
print('Тильда:',{mm:{'строк':sum(len(v) for v in per.values()),
                     'заказов':sum(sum(x[1] for x in r['w']) for r in R if r.get('src')=='tilda' and r['m']==mm)}
                 for mm,per in TIL.items()})

# ---------- АВИТО ----------
for mm,per in A_AVITO.items():
    if mm not in LIVE or not isinstance(per,dict): continue
    for prod,d in per.items():
        W=d.get('weeks') or []
        if not W: continue
        sb=d.get('selfbuy') or [0]*5; sbs=d.get('selfbuy_sum') or [0]*5
        w=blank()
        for i,x in enumerate(W[:5]):
            w[i]=[x.get('spend',0), x.get('orders',0)-sb[i], x.get('buy',0)-sb[i], x.get('buysum',0)-sbs[i],
                  x.get('views',0), 0, x.get('sum',0)-sbs[i], 0]
        R.append(dict(m=mm,prod=prod,plat='Авито',chan='Внутренняя реклама Авито',act='Авито',
            name='Объявления Авито — API (отменено %d, в пути %d)'%(sum(x.get('cancel',0) for x in W),
                                                                     sum(x.get('transit',0) for x in W)),
            link=None,tag='avito',mgr='—',w=w,est=False,src='avito',
            leads=[x.get('contacts',0) for x in W[:5]]+[0]*(5-len(W[:5])),
            msgs=d.get('msgs',0),calls=d.get('calls',0),date='—'))
        if sum(sb):
            ws=blank()
            for i in range(5): ws[i]=[0,sb[i],sb[i],0,0,0,sbs[i],0]
            R.append(dict(m=mm,prod=prod,plat='Авито',chan='Самовыкупы',act='Самовыкупы',sys=True,info=True,
                name='Самовыкупы Авито (число от Виталия) — в итог не входят',link=None,date='весь месяц',
                tag='sys',mgr='—',w=ws,est=False))
print('Авито:',{mm:[round(sum(sum(w[j] for w in r['w']) for r in R if r['m']==mm and r.get('src')=='avito'))
                    for j in (0,1,2,6)] for mm in A_AVITO if mm in LIVE})

# ---------- САЙТОВЫЕ ПОСЕВЫ: заказы только из Тильды, цифры менеджеров обнуляем ----------
# У размещений с площадкой «Сайт» заказы в таблице — оценка менеджера. Факт по сайту даёт
# только выгрузка Тильды, поэтому в строках менеджеров заказы/выкупы/выручку обнуляем
# и проставляем по совпадению метки.
SITE=[r for r in R if r['m'] in LIVE and str(r.get('plat','')).startswith('Сайт')
      and not r.get('sys') and r.get('src')!='tilda']
for r in SITE:
    for i in range(5):
        r['w'][i][1]=0; r['w'][i][2]=0; r['w'][i][3]=0; r['w'][i][6]=0
    r['src']=None
for mm,per in TIL.items():
    for c,xs in per.items():
        if not c: continue
        hit=[r for r in SITE if r['m']==mm and (r.get('camp')==c or (r.get('utmp') or {}).get('camp')==c)]
        if not hit: continue
        # заказы метки раскладываем на найденные размещения поровну
        for j,r in enumerate(hit):
            part=[x for k,x in enumerate(xs) if k%len(hit)==j]
            for x in part:
                i=widx(x['dt']); r['w'][i][1]+=1; r['w'][i][6]+=x['amt']
                if x['paid']: r['w'][i][2]+=1; r['w'][i][3]+=x['amt']
            r['src']='tilda'; r['est']=False
        # строку-дубль «метка без размещения» убираем, раз размещение нашлось
        R[:]=[z for z in R if not (z.get('src')=='tilda' and z['m']==mm and z.get('camp')==c
                                   and z.get('name','').startswith('Метка '))]

# ---------- БОТЫ КАК ИСТОЧНИК ПРОДАЖ ----------
# Заказы по меткам ботов — отдельные строки: расход 0, заказы вынимаются из органики.
BOT_OZ_CONTENT={'klevaya_karta':'km','Kiev_karta':'km'}
BOTNAME={'km':'Бот «Клёвая карта»','gd':'Бот «Зелёный доктор»','us':'Бот «Умный сад»'}
for mm in LIVE:
    cont=(A_DEEP.get('content') or {}).get(mm,{})
    acc_b={}
    for key,bot in BOT_OZ_CONTENT.items():
        d=cont.get(key)
        if not d or not d.get('o_attr'): continue
        a=acc_b.setdefault(bot,{'o':0,'s':0,'keys':[]})
        a['o']+=d['o_attr']; a['s']+=d.get('sum',0); a['keys'].append(key)
    for bot,a in acc_b.items():
        w=blank(); w[0]=[0,a['o'],0,0,0,0,a['s'],0]
        R.append(dict(m=mm,prod='atf',plat='Ozon',chan=BOTNAME[bot],act='Бот → продажи',
            name=BOTNAME[bot]+': заказы на Ozon по меткам '+', '.join(a['keys']),
            link=None,date='весь месяц',tag='utm',mgr='—',w=w,est=False,src='api',bot=bot))
    if acc_b: print('боты → продажи %s:'%mm,{k:v['o'] for k,v in acc_b.items()})

# ---------- ФАКТ ПЛОЩАДОК (из API) И ОРГАНИКА КАК ОСТАТОК ----------
FACT={}                       # (мес, площадка, продукт) -> [[заказы,выкупы,сумма] x5]
def fblank(): return [[0,0,0] for _ in range(5)]
def spread_tot(tot, wts):
    """разложить итог по неделям по долям, остаток — в самую крупную неделю"""
    out=[round(tot*x) for x in wts]; d=round(tot)-sum(out)
    if d: out[max(range(5),key=lambda i:wts[i])]+=d
    return out

for mm,d in A_OZORD.items():
    if mm not in LIVE or not isinstance(d,dict): continue
    wk=d.get('weeks') or []
    if wk:
        tq=sum(x['qty'] for x in wk) or 1
        for prod in ('atf','appl'):
            key='qty' if prod=='atf' else 'appl_qty'
            skey='sum' if prod=='atf' else 'appl_sum'
            if not sum(x.get(key,0) for x in wk): continue
            F=FACT.setdefault((mm,'Ozon',prod),fblank())
            dv=sum(x.get('delivered',0) for x in wk)
            tq2=sum(x.get(key,0) for x in wk) or 1
            for i,x in enumerate(wk):
                F[i][0]=x.get(key,0); F[i][2]=x.get(skey,0)
                F[i][1]=round(dv*x.get(key,0)/tq2) if prod=='atf' else 0
            if prod=='atf':
                dd=dv-sum(r[1] for r in F)
                if dd: F[max(range(5),key=lambda i:F[i][0])][1]+=dd
for mm,per in A_WBFUN.items():
    if mm not in LIVE or not isinstance(per,dict): continue
    for prod,d in per.items():
        if not isinstance(d,dict) or not d.get('orders'): continue
        wks=d.get('weeks') or [1,0,0,0,0]
        t=sum(wks) or 1; wts=[x/t for x in wks]
        F=FACT.setdefault((mm,'Wildberries',prod),fblank())
        o=spread_tot(d['orders'],wts); b=spread_tot(d.get('buyouts',0),wts); s=spread_tot(d.get('sum',0),wts)
        for i in range(5): F[i]=[o[i],b[i],s[i]]

BUYSUM={}
for mm,per in A_WBFUN.items():
    if mm in LIVE and isinstance(per,dict):
        for prod,d in per.items():
            if isinstance(d,dict) and d.get('buysum'): BUYSUM[(mm,'Wildberries',prod)]=d['buysum']

AVG={k:(sum(x[2] for x in v)/sum(x[0] for x in v) if sum(x[0] for x in v) else 0) for k,v in FACT.items()}
AVGBUY={k:(BUYSUM[k]/sum(x[1] for x in FACT[k]) if sum(x[1] for x in FACT[k]) else 0) for k in BUYSUM}

# органика = факт площадки − внутренняя реклама − посевы − самовыкупы
ORG={}
for (mm,pl,prod),F in sorted(FACT.items()):
    if mm not in LIVE: continue
    used=blank()
    for r in R:
        if r['m']!=mm or r.get('plat')!=pl or r.get('prod')!=prod: continue
        if r.get('act')=='Органика' or r.get('info'): continue
        for i in range(5):
            used[i][1]+=r['w'][i][1]; used[i][2]+=r['w'][i][2]
    sb=selfb.get((mm,pl),[0]*5)
    w=blank(); tot=[0,0]
    for i in range(5):
        o=max(0,round(F[i][0]-used[i][1]-sb[i]))
        b=max(0,round(F[i][1]-used[i][2]-sb[i]))
        avg=AVG.get((mm,pl,prod),0); avgb=AVGBUY.get((mm,pl,prod),avg)
        w[i]=[0,o,b,round(b*avgb),0,0,round(o*avg),0]
        tot[0]+=o; tot[1]+=b
    if sum(tot)==0: continue
    ORG[(mm,pl,prod)]=tot
    R.append(dict(m=mm,prod=prod,plat=pl,chan='Органика',act='Органика',sys=True,
        name='Остаток: факт площадки минус внутренняя реклама, посевы и самовыкупы',
        link=None,date='весь месяц',tag='sys',mgr='—',w=w,est=False))
print('ФАКТ площадок:',{('%s|%s|%s'%k):[round(sum(x[j] for x in v)) for j in range(3)] for k,v in sorted(FACT.items())})
print('ОРГАНИКА:',{('%s|%s|%s'%k):v for k,v in sorted(ORG.items())})

# ---------- ПЛАН: только из медиаплана, лист «План упрощённый» ----------
# В файле после каждого месяца идёт колонка «факт» — её НЕ читаем: факт вносится вручную
# и может расходиться с площадками. Из медиаплана берём ТОЛЬКО план.
PRODKEY={'АМИНА ТРИГГЕР ФИШ СИГНАЛ':'atf','АМИНО ТРИГГЕР ПРИКОРМКА':'atfp',
         'АМИНА ТРИГГЕР ФИШ':'atf','АППЛИКАТОР ЯРМАКОВА':'appl',
         'БОТ «ЗЕЛЁНЫЙ ДОКТОР»':'gd','БОТ «УМНЫЙ САД»':'us',
         'БОТ «КЛЁВАЯ КАРТА»':'km','БОТ «КЛЁВОЕ МЕСТО»':'km'}
SKIP={'DRR, %','CPO средний, ₽','Средняя цена регистрации, ₽','+ свободный канал','-','Итого бюджет, ₽',
      'Итого штук (упаковки / банки)','Коэффициент штук на заказ (ввод)','Средний чек, ₽ (ввод)',
      '% выкупа (ввод)','Сумма заказов, ₽','Выкупы','Выручка план, ₽','Итого заказов','Итого регистраций'}
ACT2CHAN={'ВБ внутренняя':('chan','Внутренняя реклама WB'),'Ozon внутренняя':('chan','Внутренняя реклама Ozon'),
          'Авито':('plat','Авито')}
PLANS={}
mws=openpyxl.load_workbook(SRC+'mediaplan.xlsx',data_only=True)['План упрощённый']
mrows=[list(r) for r in mws.iter_rows(min_row=1,values_only=True)]
# шапка месяцев: ищем в первых строках, план берём только из колонок с названием месяца
MONNAME={'июль':'2026-07','август':'2026-08','сентябрь':'2026-09','октябрь':'2026-10',
         'ноябрь':'2026-11','декабрь':'2026-12'}
MONCOL={}
for r in mrows[:6]:
    for j,c in enumerate(r):
        s=txt(c).lower()
        if s in MONNAME: MONCOL[j]=MONNAME[s]
    if MONCOL: break
if not MONCOL: raise SystemExit('Медиаплан: не нашёл шапку с месяцами на листе «План упрощённый».')
print('колонки медиаплана:',{k:v for k,v in sorted(MONCOL.items())})
cur=None; sec=None
for r in mrows:
    a=txt(r[0])
    if a.upper().startswith('СВОД ПО КАНАЛАМ'): break
    if a.startswith('▼'):
        t=a[1:].strip().upper(); cur=None
        for k,v in PRODKEY.items():
            if t.startswith(k): cur=v
        if cur is None:
            raise SystemExit('Медиаплан: неизвестный продукт «%s». Добавь его в PRODKEY, иначе план по нему потеряется.'%a[1:].strip())
        sec=None; continue
    if a.startswith('──'): sec=a; continue
    if not cur: continue
    for ci,mm in MONCOL.items():
        P=PLANS.setdefault(mm,{'prod':{},'plat':{},'chan':{},'act':{}})
        d=P['prod'].setdefault(cur,{})
        v=num(r[ci]) if ci<len(r) else 0
        if a=='Итого бюджет, ₽': d['b']=round(v)
        elif a=='Итого заказов': d['o']=round(v)
        elif a=='Итого регистраций': d['r']=round(v)
        elif a=='Выкупы': d['buy']=round(v)
        elif a=='Выручка план, ₽': d['rev']=round(v)
        elif sec and 'Бюджет по каналам' in sec and a and a not in SKIP:
            if v: P['act'].setdefault(a,{}).setdefault(cur,{})['b']=round(v)
        elif sec and ('Заказы по каналам' in sec or 'Регистрации (расч' in sec) and a and a not in SKIP:
            if v: P['act'].setdefault(a,{}).setdefault(cur,{})['o']=round(v)
ACTMAP={'Инфлюенс-маркетинг':'Инфлюенс-посев','ВБ внутренняя':'Внутренняя реклама','Ozon внутренняя':'Внутренняя реклама',
        'Авито':'Авито','SEO':'SEO-продвижение','Конкурс репостов в вк':'Конкурс репостов в ВК',
        'Конкурс репостов':'Конкурс репостов в ВК','Телеграм АДС':'Телеграм ADS',
        'YouTube-канал':'Свой YouTube-канал','Бартерная реклама':'Бартер','Таргет ВК':'Таргет ВК',
        'Контекст (Директ)':'Контекст'}
for mm,P in PLANS.items():
    for act,per in P['act'].items():
        kind,name=ACT2CHAN.get(act,(None,None))
        if act=='Инфлюенс-маркетинг' or not kind: continue
        for prod,vv in per.items():
            if kind=='chan': P['chan'].setdefault(prod,{})[name]=dict(vv)
            elif kind=='plat': P['plat'].setdefault(prod,{})[name]=dict(vv)
for mm in list(PLANS):
    P=PLANS[mm]
    P['prod']={k:v for k,v in P['prod'].items() if v.get('b') or v.get('o') or v.get('r')}
    P['act']={k:v for k,v in P['act'].items() if v}
    if not P['prod'] and not P['act']: del PLANS[mm]
print('план по месяцам:',{mm:{'бюджет':sum(v.get('b',0) for v in P['prod'].values()),
                             'заказы':sum(v.get('o',0) for v in P['prod'].values()),
                             'регистрации':sum(v.get('r',0) for v in P['prod'].values())}
                          for mm,P in sorted(PLANS.items())})

# ---------- СКЛЕЙКА С ЗАМОРОЖЕННОЙ ИСТОРИЕЙ И СБОРКА META ----------
FR=FROZ
ALL = FR['records'] + R

# процент выкупа (когортно): закрытые месяцы — из снимка, живые — считаем
buyrate=dict(FR.get('buyrate') or {})
buyrate['open']=[TODAY.strftime('%Y-%m')]
for (mm,pl,prod),F in FACT.items():
    o=sum(x[0] for x in F); b=sum(x[1] for x in F)
    buyrate.setdefault(mm,{}).setdefault(pl,{})[prod]=(round(b/o*100) if o and mm!=TODAY.strftime('%Y-%m') else None)

# рейтинг по неделям: закрытые — из снимка, живые — из снимков rating.json
rating={k:v for k,v in (FR.get('rating') or {}).items()}
for mm,per in (RATE or {}).items():
    if mm in LIVE: rating.setdefault(mm,{}).update(per)
HIST=(A_RATE.get('history') or [])
ratingLast={}
for h in HIST:
    try: dt=datetime.datetime.strptime(h['date'],'%Y-%m-%d')
    except Exception: continue
    mm=dt.strftime('%Y-%m')
    if mm not in MONTHS: continue
    wi=widx(dt)
    for pl,per in h.items():
        if pl=='date' or not isinstance(per,dict): continue
        for prod,v in per.items():
            rating.setdefault(mm,{}).setdefault(pl,{}).setdefault(prod,[None]*5)[wi]=v
            ratingLast[pl]=dict(v=v,when=MONTHS[mm]['weeks'][wi],prod=prod)

# заказы по меткам Ozon для вкладки «Атрибуция публикаций»
WHO={}
for r in ALL:
    if r.get('ckey') and r.get('plat')=='Ozon' and not r.get('sys'):
        WHO.setdefault((r['m'],tuple(r['ckey'])),[]).append(r['name'])
ozutm={}
for mm,per in A_UTM.items():
    if mm not in MONTHS or not isinstance(per,dict): continue
    lst=[]
    for camp,d in per.items():
        if not isinstance(d,dict): continue
        lst.append(dict(camp=camp,s=d.get('s',0),c=d.get('c',0),cart=d.get('cart',0),
                        o_sess=d.get('o_sess',0),o_attr=d.get('o_attr',0),sum=d.get('sum',0),
                        who=WHO.get((mm,('camp',camp)),[])))
    lst.sort(key=lambda x:-x['sum']); ozutm[mm]=lst
for mm,v in (FR.get('ozutm') or {}).items(): ozutm.setdefault(mm,v)

utmhit={'%s|%s|%s'%k:v for k,v in UTMHIT.items()}
utmhit.update(FR.get('utmhit') or {})

pend={'%s|%s'%k:v for k,v in PEND.items()}
pend.update(FR.get('pend') or {})

fact={'%s|%s|%s'%k:[round(sum(x[j] for x in v)) for j in range(3)] for k,v in FACT.items()}
fact.update(FR.get('fact') or {})

sb_all={}
for r in ALL:
    if r.get('act')=='Самовыкупы': sb_all[r['plat']]=sb_all.get(r['plat'],0)+sum(x[1] for x in r['w'])

def _mt(*fs):
    t=max((os.path.getmtime(f) for f in fs if os.path.exists(f)), default=0)
    return datetime.datetime.fromtimestamp(t).strftime('%Y-%m-%d %H:%M') if t else '—'
srcupd=[
 dict(n='Ozon — API', ok=30, w='заказы, реклама, UTM-метки, рейтинг',
      d=_mt(API+'ozon_orders.json',API+'ozon_ads.json',API+'ozon_utm.json',API+'ozon_utm_deep.json')),
 dict(n='Wildberries — API', ok=30, w='воронка продаж, реклама, рейтинг', d=_mt(API+'wb_funnel.json',API+'wb_ads.json')),
 dict(n='Авито — API', ok=30, w='заказы, показы, контакты, расход', d=_mt(API+'avito.json')),
 dict(n='Таблицы менеджеров', ok=72, w='размещения, расход, охваты, метки',
      d=_mt(SRC+'elena.xlsx',SRC+'maria.xlsm',SRC+'anastasia.xlsx')),
 dict(n='Сайт — выгрузка Тильды', ok=72, w='заказы и выкупы с trigger-fish.ru', d=_mt(SRC+'tilda_orders.csv')),
 dict(n='Медиаплан', ok=960, w='план по бюджету, заказам, выкупам, выручке', d=_mt(SRC+'mediaplan.xlsx')),
]
avito_meta={mm:{'msgs':(per.get('atf') or {}).get('msgs',0),'calls':(per.get('atf') or {}).get('calls',0),
                'contacts':(per.get('atf') or {}).get('msgs',0)+(per.get('atf') or {}).get('calls',0)}
            for mm,per in A_AVITO.items() if mm in MONTHS and isinstance(per,dict)}

meta=dict(months=MONTHS, buyrate=buyrate,
          ratingApi=dict(date=A_RATE.get('date'),month=A_RATE.get('month'),week=A_RATE.get('week_index'),n=len(HIST)),
          ratingContent=A_RATE.get('content') or {},
          plans={mm:v for mm,v in PLANS.items() if mm in MONTHS},
          rating=rating, ratingLast=ratingLast, pend=pend,
          selfbuy={'wb':round(sb_all.get('Wildberries',0)),'oz':round(sb_all.get('Ozon',0))},
          avito=avito_meta, updated=TODAY.strftime('%d.%m.%Y'), srcupd=srcupd, fact=fact,
          actmap=ACTMAP, ozutm=ozutm, utmhit=utmhit,
          apisrc=dict(wb_funnel=sorted(k for k in A_WBFUN if k in MONTHS),
                      wb_ads=sorted(k for k in A_WBADS if k in MONTHS),
                      oz_ads=sorted(k for k in A_OZADS if k in MONTHS),
                      oz_utm=sorted(k for k in A_UTM if k in MONTHS)),
          check=dict(el=[round(x) for x in el.get('2026-08',[0,0,0,0])],
                     ma=[round(x) for x in ma.get('2026-08',[0,0,0,0])]))
json.dump(dict(records=ALL,meta=meta),open('records.json','w'),ensure_ascii=False)
print()
print('ЗАПИСЕЙ ВСЕГО: %d (заморожено %d + пересчитано %d)'%(len(ALL),len(FR['records']),len(R)))
print('месяцы:',dict(_Cnt(r['m'] for r in ALL)))
def agg(rs,j): return sum(sum(w[j] for w in r['w']) for r in rs)
for mm in LIVE:
    S=[r for r in ALL if r['m']==mm and not r.get('info')]
    print('%s: расход %9.0f | заказы %5d | выкупы %5d | сумма %10.0f'%(mm,agg(S,0),agg(S,1),agg(S,2),agg(S,6)))
